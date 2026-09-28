"""Read-only Excel extraction and validation; never modifies the workbook."""
import sys,json,re,hashlib,datetime,unicodedata
from urllib.parse import urlsplit
import openpyxl

def text(v): return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]','',str(v if v is not None else '')).strip()
def header(v): return ' '.join(unicodedata.normalize('NFKC',text(v)).lower().split())
def code(v):
    if isinstance(v,float) and v.is_integer(): v=int(v)
    s=text(v)
    if not re.fullmatch(r'[0-9]+',s) or int(s)==0: raise ValueError('Invalid numeric companyCode')
    return s

def rows(sheet,required):
    found=False
    for row in sheet.iter_rows(values_only=True):
        if not found:
            names=[header(v) for v in row]
            if not all(header(k) in names for k in required): continue
            nonempty=[n for n in names if n]
            if len(set(nonempty))!=len(nonempty): raise ValueError('Duplicate headers: '+sheet.title)
            found=True
            continue
        if any(v is not None for v in row): yield dict(zip(names,row))
    if not found: raise ValueError('Missing columns: '+sheet.title)

def extract(path,version):
    w=openpyxl.load_workbook(path,read_only=True,data_only=True)
    assert set(w.sheetnames)==set(['Overview','Universe','Ranking','Red Flags','Sources','Top 10'])
    asof=None
    for r in w['Overview'].values:
        if len(r)>1 and header(r[0])=='snapshot date':
            d=r[1]; asof=d.date().isoformat() if isinstance(d,datetime.datetime) else datetime.datetime.strptime(str(d),'%d-%b-%Y').date().isoformat()
    if not asof: raise ValueError('Missing snapshot date')
    fields={'companyName':'Company','promoterController':'Promoter / Controller','mainStrength':'Main Strength','primaryFlag':'Primary Red / Yellow Flag','regulatoryTaxFinding':'Reg/Tax Flag','relatedPartyFinding':'Heavy RPT','promoterPayFinding':'Excess Pay','capitalAllocationFinding':'Unrelated Allocation','promoterActivity':'Promoter Activity','evidenceNote':'Evidence Note'}
    factors=[('regulatory','Regulatory',.25,'regulatoryScore'),('remuneration','Remuneration',.15,'remunerationScore'),('relatedPartyTransactions','RPT',.20,'rptScore'),('capitalAllocation','Capital Allocation',.20,'capitalAllocationScore'),('skinInTheGame','Skin',.10,'skinInTheGameScore'),('disclosure','Disclosure',.10,'disclosureScore')]
    required=['Company Code','Overall Rank','Weighted Score','Risk',*fields.values(),*[f[1] for f in factors],'Source 1','Source 2','Source 3']
    ranking=list(rows(w['Ranking'],required)); audits={}; sources={}
    def add_source(c,u,n,note):
        u=text(u)
        if not u:return
        parsed=urlsplit(u)
        if parsed.scheme not in ('http','https') or not parsed.hostname or parsed.username or parsed.password: raise ValueError('Invalid source URL')
        if c not in audits: raise ValueError('Source code absent from Ranking')
        if not float(n).is_integer() or int(n)<1: raise ValueError('Invalid source number')
        k=c+'\n'+u
        if k in sources:
            old=sources[k]; old['sourceNumber']=min(old['sourceNumber'],int(n))
            if note and note not in old['auditNote'].split('\n\n'):old['auditNote']+= ('\n\n' if old['auditNote'] else '')+note
        else:sources[k]={'companyCode':c,'url':u,'sourceNumber':int(n),'sourceLabel':parsed.hostname.removeprefix('www.'),'sourceType':'PDF' if parsed.path.lower().endswith('.pdf') else 'WEB','auditNote':note}
    for r in ranking:
        c=code(r['company code'])
        if c in audits: raise ValueError('Duplicate companyCode '+c)
        a={k:text(r[header(v)]) for k,v in fields.items()};a['companyCode']=c;a['scores']={}
        for key,col,weight,scalar in factors:
            n=float(r[header(col)])
            if not n.is_integer() or not 1<=n<=10:raise ValueError('Invalid score '+c)
            a['scores'][key]=a[scalar]=int(n)
        score=float(r['weighted score']);calc=round(sum(a['scores'][key]*weight for key,_,weight,_ in factors),2)
        if not 1<=score<=10 or abs(score-calc)>.000001:raise ValueError('Weighted score mismatch '+c)
        rank=float(r['overall rank'])
        if not rank.is_integer() or rank<1:raise ValueError('Invalid rank')
        risk=text(r['risk']).upper()
        if risk not in ('LOW','MEDIUM','HIGH'):raise ValueError('Invalid risk')
        a.update(overallRank=int(rank),promoterScore=score,riskCategory=risk,auditVersion=version,auditAsOfDate=asof)
        audits[c]=a
    source_rows=list(rows(w['Sources'],['Company Code','Source #','URL','Audit Note']))
    for r in source_rows:add_source(code(r['company code']),r['url'],r['source #'],text(r['audit note']))
    for r in ranking:
        for n in range(1,4):add_source(code(r['company code']),r['source '+str(n)],n,'')
    top=[]
    for r in rows(w['Top 10'],['Rank','Company Code','Promoter Score (1-10)','Risk Category']):
        c=code(r['company code']);rank=int(r['rank']);a=audits[c]
        if float(r['rank'])!=rank or not 1<=rank<=10 or a['promoterScore']!=float(r['promoter score (1-10)']) or a['riskCategory']!=text(r['risk category']).upper():raise ValueError('Invalid Top 10')
        top.append({'rank':rank,'companyCode':c})
    if len(top)!=10 or len({t['companyCode'] for t in top})!=10 or len({t['rank'] for t in top})!=10:raise ValueError('Duplicate/incomplete Top 10')
    w.close()
    return dict(audits=list(audits.values()),sources=list(sources.values()),top10=sorted(top,key=lambda t:t['rank']),auditVersion=version,auditAsOfDate=asof,workbookHash=hashlib.sha256(open(path,'rb').read()).hexdigest(),methodology={'weights':{key:weight for key,_,weight,_ in factors},'higherIsBetter':True},workbookRowsRead={'Ranking':len(ranking),'Sources':len(source_rows),'Top 10':len(top)})
if __name__=='__main__':
    result=extract(sys.argv[1],sys.argv[3])
    with open(sys.argv[2],'w') as f:json.dump(result,f,ensure_ascii=False)
    print(json.dumps({'audits':len(result['audits']),'sources':len(result['sources']),'date':result['auditAsOfDate']}))
